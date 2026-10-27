import sys
import os
import subprocess

def switch_to_venv():
    """
    現在の実行環境が .venv でなければ、2個上にある .venv の Python で
    自分自身を自動的に再起動（リブート）する
    """
    # このスクリプトがある場所を基準にする
    current_dir = os.path.dirname(os.path.abspath(__file__))
    
    # 2個上のディレクトリにある .venv の Python のパスを正確に割り出す
    # （※もし1個上の階層なら '../.venv/Scripts/python.exe' に変更してください）
    target_python = os.path.abspath(os.path.join(current_dir, "../../.venv/Scripts/python.exe"))
    
    # すでに .venv の Python で動いている場合はそのまま進む
    if os.path.normcase(os.path.abspath(sys.executable)) == os.path.normcase(target_python):
        return

    # .venv の Python が存在する場合、そちらに切り替えて自分を再実行する
    if os.path.exists(target_python):
        print(f"[自動切替] .venv の環境に切り替えます...")
        try:
            # 自分自身を .venv の Python でもう一度起動し、引数も引き継ぐ
            result = subprocess.run([target_python] + sys.argv)
            # 再起動したプロセスの終了コードを引き継いで終了
            sys.exit(result.returncode)
        except Exception as e:
            print(f"環境の切り替えに失敗しました: {e}")
            sys.exit(1)
    else:
        print(f"【警告】指定された仮想環境が見つかりません: {target_python}")
        print("システム側の Python で続行します...")

# 実行の最最初に必ずチェックさせる
switch_to_venv()


import os
import re
import sys
import traceback
import pandas as pd
from openpyxl import load_workbook

# 空のテンプレートファイルのパス
TEMPLATE_PATH = 'template.xlsm'

def extract_groups(filename):
    """
    ファイル名を分割し、3段階のグループを抽出する。
    """
    parts = re.split(r'[_ ]', filename)
    group1 = parts[0] if len(parts) > 0 else filename
    group2 = parts[1] if len(parts) > 1 else ''
    group3 = parts[2] if len(parts) > 2 else ''
    return group1, group2, group3

def extract_number(part):
    """
    部分文字列から最初に見つかった数値を抽出して返す。
    """
    match = re.search(r'\d+', part)
    return int(match.group()) if match else float('inf')

def custom_sort_key(col):
    """
    ソートキーを生成する。
    """
    def key_fn(value):
        if isinstance(value, str):
            match = re.match(r'(\D*)(\d*)', value)
            prefix = match.group(1).lower()
            number = int(match.group(2)) if match.group(2).isdigit() else float('inf')
            return (prefix, number)
        return value
    return col.map(key_fn)

def save_filenames_to_excel():
    print("[1/2] ファイル名の一覧化処理を開始します...")
    all_files = os.listdir()
    
    # 対象ファイルをフィルタリング
    valid_files = [os.path.splitext(f)[0] for f in all_files if not f.endswith(('.py', '.txt', '.xlsx', '.xlsm'))]
    
    if not valid_files:
        print("  -> 【注意】対象となるファイル（.py, .txt, .xlsx, .xlsm 以外）が見つかりませんでした。")
        return

    data = []
    for filename in valid_files:
        group1, group2, group3 = extract_groups(filename)
        number1 = extract_number(group1)
        number2 = extract_number(group2)
        number3 = extract_number(group3)
        data.append((filename, group1, group2, group3, number1, number2, number3))
    
    df = pd.DataFrame(data, columns=['File Name', 'Group1', 'Group2', 'Group3', 'Number1', 'Number2', 'Number3'])
    
    df = df.sort_values(
        by=['Group1', 'Number1', 'Group2', 'Number2', 'Group3', 'Number3'],
        key=custom_sort_key
    )
    
    excel_filename = 'non_py_txt_xlsx_xlsm_files.xlsx'
    df[['File Name']].to_excel(excel_filename, index=False)
    print(f"  -> 成功: '{excel_filename}' を保存しました。")

def save_txt_to_single_excel():
    print("[2/2] テキストファイルの結合とテンプレートへの書き込みを開始します...")
    
    txt_files = [f for f in os.listdir() if f.endswith('.txt')]
    
    if not txt_files:
        print("  -> 【エラー】フォルダ内に '.txt' ファイルが1つも見つかりません！")
        return

    print(f"  -> 見つかったテキストファイル: {txt_files}")
    
    combined_df = pd.DataFrame()
    
    for txt_file in txt_files:
        print(f"  -> 読み込み中: {txt_file}")
        df = pd.read_csv(txt_file, delimiter='\t')
        combined_df = pd.concat([combined_df, df], ignore_index=True)
    
    combined_filename = 'combined_txt_files.xlsx'
    combined_df.to_excel(combined_filename, index=False)
    print(f"  -> 成功: '{combined_filename}' を一時保存しました。")
    
    if not os.path.exists(TEMPLATE_PATH):
        print(f"  -> 【エラー】テンプレートファイル '{TEMPLATE_PATH}' がカレントディレクトリに存在しません！")
        return
    
    print(f"  -> テンプレート '{TEMPLATE_PATH}' を読み込んでいます...")
    workbook = load_workbook(TEMPLATE_PATH, keep_vba=True)
    sheet = workbook.active

    # 既存データをクリア
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.value = None

    # データを書き込み
    print("  -> データをテンプレートに書き込んでいます...")
    for i, row in enumerate(combined_df.values, start=2):
        for j, value in enumerate(row, start=1):
            sheet.cell(row=i, column=j, value=value)
    
    txt_file_name = txt_files[0]
    base_name = txt_file_name.split(' ')[0]
    
    xlsm_filename = f"{base_name}_spectre.xlsm"
    
    workbook.save(xlsm_filename)
    print(f"  -> 成功: '{xlsm_filename}' を保存しました！")

def main():
    save_filenames_to_excel()
    print("-" * 30)
    save_txt_to_single_excel()

if __name__ == '__main__':
    try:
        main()
        print("\n" + "="*50)
        print(" すべての処理が正常に完了しました！")
        print("="*50)
    except Exception:
        print("\n" + "!"*60)
        print(" 【重大なエラーが発生しました。以下の詳細を確認してください】")
        print("!"*60)
        # どこの何行目でエラーが起きたかを完璧に出力する
        traceback.print_exc()
        print("!"*60)
    finally:
        # どんな結果であれ、Enterを押すまで画面を絶対に閉じさせない
        input("\n終了するには Enter キーを押してください...")